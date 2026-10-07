#!/usr/bin/env python3
"""Build an editable PowerPoint from EduPPT Creator JSON input."""
import argparse
import json
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(23, 54, 93)
BLUE = RGBColor(47, 102, 144)
PALE = RGBColor(238, 245, 250)
DARK = RGBColor(35, 43, 52)
GREY = RGBColor(95, 105, 115)
WHITE = RGBColor(255, 255, 255)

def rgb_from_hex(value, default=BLUE):
    try:
        value = value.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    except Exception:
        return default

def set_text(shape, text, size, color, bold=False, font="Aptos"):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = str(text)
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return tf

def add_text(slide, x, y, w, h, text, size, color=DARK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(shape, text, size, color, bold)
    return shape

def add_bullets(slide, x, y, w, h, bullets, size=22):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.04)
    for i, item in enumerate(bullets or []):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(item)
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(14)
        p.bullet = True
    return shape

def add_header(slide, title, accent=BLUE):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, 0.65, 0.42, 12.0, 0.72, title, 29, NAVY, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.25), Inches(1.35), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()

def add_footer(slide, number, source_note=""):
    add_text(slide, 0.65, 7.15, 10.8, 0.2, source_note[:160], 8, GREY)
    add_text(slide, 12.0, 7.10, 0.6, 0.25, str(number), 10, GREY)

def add_cover(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.15), Inches(0.12), Inches(4.7))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(91, 192, 190)
    accent.line.fill.background()
    add_text(slide, 1.3, 1.55, 10.8, 1.65, data["title"], 34, WHITE, True)
    subtitle = data.get("subtitle", "")
    if subtitle:
        add_text(slide, 1.33, 3.35, 10.5, 0.8, subtitle, 19, RGBColor(220, 231, 240))
    meta = " · ".join(x for x in [data.get("subject", ""), data.get("grade", "")] if x)
    if meta:
        add_text(slide, 1.33, 5.45, 10.0, 0.45, meta, 15, RGBColor(180, 210, 225), True)
    return slide

def add_standard_slide(prs, item, number, source_note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = rgb_from_hex(item.get("accent", "#2F6690"))
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_header(slide, item.get("title", f"Slide {number}"), accent)
    layout = item.get("layout", "bullets")
    if layout == "two_column":
        for x, title_key, bullets_key in [
            (0.75, "left_title", "left_bullets"),
            (6.85, "right_title", "right_bullets")
        ]:
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(5.65), Inches(4.95))
            card.fill.solid()
            card.fill.fore_color.rgb = PALE
            card.line.color.rgb = RGBColor(220, 230, 238)
            add_text(slide, x+0.25, 1.9, 5.1, 0.5, item.get(title_key, ""), 20, accent, True)
            add_bullets(slide, x+0.25, 2.55, 5.0, 3.7, item.get(bullets_key, []), 18)
    elif layout == "timeline":
        bullets = item.get("bullets", [])
        y = 1.8
        for i, value in enumerate(bullets):
            cy = y + i * min(0.95, 4.4/max(1, len(bullets)))
            circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(cy), Inches(0.38), Inches(0.38))
            circ.fill.solid()
            circ.fill.fore_color.rgb = accent
            circ.line.fill.background()
            add_text(slide, 1.55, cy-0.02, 10.8, 0.65, value, 18, DARK)
    elif layout == "closing":
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.75), Inches(4.95))
        panel.fill.solid()
        panel.fill.fore_color.rgb = PALE
        panel.line.fill.background()
        add_bullets(slide, 1.15, 2.0, 10.9, 4.1, item.get("bullets", []), 21)
    else:
        add_bullets(slide, 0.9, 1.65, 11.7, 4.95, item.get("bullets", []), 21)
    note = item.get("note", "") or source_note
    add_footer(slide, number, note)
    return slide

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", help="Path to lesson JSON")
    parser.add_argument("--output", default="output/lesson.pptx", help="Output PPTX path")
    args = parser.parse_args()
    with open(args.input_json, "r", encoding="utf-8") as f:
        data = json.load(f)
    if not data.get("title") or not isinstance(data.get("slides"), list) or not data["slides"]:
        raise SystemExit("Input JSON must contain a title and a non-empty slides array.")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = data["title"]
    prs.core_properties.subject = "EduPPT Creator - THCS History and Geography"
    add_cover(prs, data)
    for i, item in enumerate(data["slides"], start=2):
        if not item.get("title"):
            raise SystemExit(f"Slide {i} is missing a title.")
        add_standard_slide(prs, item, i, data.get("source_note", ""))
    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    prs.save(args.output)
    # Reopen the saved file as a basic integrity check.
    check = Presentation(args.output)
    if len(check.slides) != len(data["slides"]) + 1:
        raise SystemExit("Validation failed: slide count does not match input.")
    print(f"Created: {os.path.abspath(args.output)}")
    print(f"Slides: {len(check.slides)} (including cover)")
    print("Basic validation: passed (file reopened and slide count checked).")
    print("Note: visual rendering/overflow checks require a presentation renderer.")

if __name__ == "__main__":
    main()
